"""Builds the mixed-format evaluation corpus and the answers to its table questions.

The Markdown report and the TXT policy are checked in as-is. The DOCX handbook
and the shipments CSV are generated here, deterministically, so the expected
answers to the aggregation questions are computed from the same rows the
system ingests rather than typed in by hand.
"""

import csv
import io
import random
from pathlib import Path

DOCUMENTS_DIR = Path(__file__).resolve().parent / "documents"

ORIGINS = ["Rotterdam", "Antwerp", "Lyon", "Poznan", "Gdansk"]
DESTINATIONS = ["Milan", "Paris", "Berlin", "Madrid", "Vienna", "Prague"]


def shipment_rows() -> list[dict]:
    rng = random.Random(7)
    rows = []
    for number in range(1, 121):
        weight = rng.randint(200, 24000)
        rows.append(
            {
                "shipment_id": f"SH-{number:04d}",
                "month": f"2025-{rng.randint(1, 12):02d}",
                "origin": rng.choice(ORIGINS),
                "destination": rng.choice(DESTINATIONS),
                "weight_kg": weight,
                "cost_eur": round(weight * 0.08 + rng.randint(100, 900), 2),
                "on_time": "yes" if rng.random() < 0.85 else "no",
            }
        )
    return rows


def shipments_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(shipment_rows()[0]))
    writer.writeheader()
    writer.writerows(shipment_rows())
    return buffer.getvalue().encode()


def handbook_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Northwind Employee Handbook (excerpts)", level=1)
    document.add_heading("Annual leave", level=2)
    document.add_paragraph(
        "Full-time employees receive 25 days of annual leave per calendar year, in "
        "addition to public holidays. Up to 5 unused days can be carried over into "
        "the next year, but carried-over days must be taken by 31 March or they "
        "are lost. Leave requests should be made at least two weeks in advance."
    )
    document.add_heading("Remote work", level=2)
    document.add_paragraph(
        "Office-based staff may work from home up to two days per week, agreed "
        "with their manager. Core hours are 10:00 to 15:00, during which everyone "
        "is expected to be reachable, wherever they are working."
    )
    document.add_heading("Probation", level=2)
    document.add_paragraph(
        "New employees serve a probation period of six months, during which the "
        "notice period on either side is two weeks. A review meeting is held "
        "before the end of probation."
    )
    document.add_heading("Equipment", level=2)
    document.add_paragraph(
        "Every office employee receives a laptop, which is replaced every three "
        "years or sooner if it fails. Phones are provided only for roles that "
        "require on-call availability."
    )
    document.add_heading("Benefits", level=2)
    table = document.add_table(rows=4, cols=2)
    for row, (name, detail) in zip(
        table.rows,
        [
            ("Benefit", "Detail"),
            ("Gym subsidy", "30 EUR per month towards any gym membership"),
            ("Bike lease", "Tax-efficient bicycle lease through JobBike"),
            ("Pension", "Employer contributes 5% of base salary"),
        ],
        strict=True,
    ):
        row.cells[0].text = name
        row.cells[1].text = detail
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def corpus() -> dict[str, bytes]:
    """Filename to file content, in upload order."""
    return {
        "operations_report_2025.md": (
            DOCUMENTS_DIR / "operations_report_2025.md"
        ).read_bytes(),
        "travel_expense_policy.txt": (
            DOCUMENTS_DIR / "travel_expense_policy.txt"
        ).read_bytes(),
        "employee_handbook.docx": handbook_docx(),
        "shipments_2025.csv": shipments_csv(),
    }


def table_answers() -> dict[str, float | str]:
    rows = shipment_rows()
    rotterdam = [r["weight_kg"] for r in rows if r["origin"] == "Rotterdam"]
    return {
        "total_cost": round(sum(r["cost_eur"] for r in rows), 2),
        "avg_weight_rotterdam": round(sum(rotterdam) / len(rotterdam), 1),
        "count_milan": sum(r["destination"] == "Milan" for r in rows),
        "on_time_pct": round(
            100 * sum(r["on_time"] == "yes" for r in rows) / len(rows), 1
        ),
        "max_cost_id": max(rows, key=lambda r: r["cost_eur"])["shipment_id"],
        "sh0042_destination": rows[41]["destination"],
    }
