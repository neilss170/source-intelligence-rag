"""Uploads, background indexing, statuses and deletion (milestone 2)."""

import io
import uuid
import zipfile
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document as WordDocument
from sqlalchemy import select, text, update
from sqlalchemy.orm import Session

from app import worker
from app.core.config import get_settings
from app.db.models import Document, DocumentTable, Passage
from app.db.session import get_sessionmaker
from app.ingest import pipeline
from app.ingest.pipeline import index_document

REPORT = (
    b"# Annual report\n\n## Safety\n\nLost-time incidents fell to 7 in 2025. "
    b"Forklift training is mandatory every 18 months.\n\n## Fleet\n\n"
    b"The company operates 214 trucks."
)
CSV = b"shipment_id,destination,cost_eur\nSH-1,Milan,100.5\nSH-2,Paris,200\n"


@pytest.fixture
def user(make_client, sign_up, ingestion):
    client = make_client()
    sign_up(client, "neil@example.com")
    return client


def upload(client, *files):
    return client.post(
        "/api/documents",
        files=[
            ("files", (name, content, "application/octet-stream"))
            for name, content in files
        ],
    )


def test_upload_is_accepted_then_indexed_into_passages(user, db) -> None:
    response = upload(user, ("report.md", REPORT))

    assert response.status_code == 202
    accepted = response.json()[0]
    assert accepted["status"] == "uploaded" and accepted["kind"] == ".md"
    document = user.get(f"/api/documents/{accepted['id']}").json()
    assert document["status"] == "ready"
    assert document["passage_count"] == 2
    assert document["error"] is None

    with Session(db) as session:
        passages = session.scalars(select(Passage).order_by(Passage.ordinal)).all()
        keyword_hit = session.execute(
            text(
                "SELECT title FROM passages "
                "WHERE tsv @@ plainto_tsquery('english', 'forklift')"
            )
        ).scalar_one()
    assert [p.title for p in passages] == [
        "report > Annual report > Safety",
        "report > Annual report > Fleet",
    ]
    assert len(passages[0].embedding) == 384
    assert keyword_hit == "report > Annual report > Safety"


def test_spreadsheets_also_store_their_rows_for_sql(user, db) -> None:
    upload(user, ("shipments.csv", CSV))

    with Session(db) as session:
        table = session.scalar(select(DocumentTable))
    assert table.row_count == 2
    assert table.rows[0] == ["shipment_id", "destination", "cost_eur"]


def test_word_documents_are_indexed(user) -> None:
    word = WordDocument()
    word.add_heading("Handbook", level=1)
    word.add_paragraph("Employees receive 25 days of annual leave.")
    buffer = io.BytesIO()
    word.save(buffer)

    document = upload(user, ("handbook.docx", buffer.getvalue())).json()[0]

    assert user.get(f"/api/documents/{document['id']}").json()["status"] == "ready"


def test_the_same_bytes_are_not_stored_twice(user, db) -> None:
    first = upload(user, ("report.md", REPORT)).json()[0]
    again = upload(user, ("renamed.md", REPORT)).json()[0]
    both_at_once = upload(user, ("a.csv", CSV), ("b.csv", CSV)).json()

    assert again["id"] == first["id"] and again["duplicate"] is True
    assert [d["duplicate"] for d in both_at_once] == [False, True]
    assert both_at_once[0]["id"] == both_at_once[1]["id"]
    with db.connect() as connection:
        assert connection.execute(text("SELECT count(*) FROM documents")).scalar() == 2


@pytest.mark.parametrize(
    ("name", "content", "code"),
    [
        ("notes.pdf", b"just text pretending", "unsupported_type"),
        ("tool.exe", b"MZ\x90\x00binary", "unsupported_type"),
        ("binary.txt", b"abc\x00\x01\x02", "unsupported_type"),
        ("latin1.txt", "café".encode("latin-1"), "unsupported_encoding"),
    ],
)
def test_files_are_judged_by_content_not_name(user, name, content, code) -> None:
    response = upload(user, (name, content))

    assert response.status_code == 415
    assert response.json()["error"]["code"] == code


def test_an_archive_named_docx_is_refused(user) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("payload.sh", "rm -rf /")

    response = upload(user, ("innocent.docx", buffer.getvalue()))

    assert response.status_code == 415


def test_size_limits(user, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "max_file_size_mb", 1)
    big = upload(user, ("big.txt", b"a" * (1024 * 1024 + 1)))
    monkeypatch.setattr(settings, "max_workspace_mb", 1)
    upload(user, ("one.txt", b"b" * 600_000))
    full = upload(user, ("two.txt", b"c" * 600_000))
    empty = upload(user, ("empty.txt", b""))

    assert big.status_code == 413 and big.json()["error"]["code"] == "file_too_large"
    assert full.status_code == 413 and full.json()["error"]["code"] == "workspace_full"
    assert empty.status_code == 422


def test_oversized_bodies_are_refused_before_they_are_read(user) -> None:
    response = user.post(
        "/api/documents",
        content=b"x",
        headers={"Content-Length": str(10**12), "Content-Type": "multipart/form-data"},
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "request_too_large"


def test_an_unreadable_file_fails_with_a_reason_and_no_passages(user, db) -> None:
    document = upload(user, ("broken.json", b'{"not": "closed"')).json()[0]

    result = user.get(f"/api/documents/{document['id']}").json()

    assert result["status"] == "failed"
    assert "could not be read" in result["error"]
    with db.connect() as connection:
        assert connection.execute(text("SELECT count(*) FROM passages")).scalar() == 0


def test_delete_removes_every_trace(user, ingestion, db) -> None:
    document = upload(user, ("shipments.csv", CSV)).json()[0]
    with Session(db) as session:
        key = session.get(Document, uuid.UUID(document["id"])).storage_key
    assert ingestion.storage.exists(key)

    assert user.delete(f"/api/documents/{document['id']}").status_code == 204

    assert user.get(f"/api/documents/{document['id']}").status_code == 404
    assert not ingestion.storage.exists(key)
    with db.connect() as connection:
        for table in ("documents", "passages", "document_tables"):
            count = connection.execute(text(f"SELECT count(*) FROM {table}")).scalar()
            assert count == 0, table


def test_a_document_deleted_while_indexing_stays_deleted(
    user, ingestion, monkeypatch, db
):
    ingestion.fail = True  # accept the upload but don't index it yet
    document_id = uuid.UUID(upload(user, ("report.md", REPORT)).json()[0]["id"])
    real_build = pipeline._build

    def build_then_delete(document, storage, settings):
        built = real_build(document, storage, settings)
        with db.begin() as connection:
            connection.execute(
                text("DELETE FROM documents WHERE id = :id"), {"id": document_id}
            )
        return built

    monkeypatch.setattr(pipeline, "_build", build_then_delete)
    with get_sessionmaker()() as session:
        status = index_document(session, document_id, ingestion.storage)

    assert status == "deleted"
    with db.connect() as connection:
        assert connection.execute(text("SELECT count(*) FROM passages")).scalar() == 0


def test_documents_the_queue_lost_are_swept_up(user, ingestion, db) -> None:
    ingestion.fail = True
    response = upload(user, ("report.md", REPORT))
    lost = uuid.UUID(response.json()[0]["id"])

    assert response.status_code == 202
    assert worker.stale_document_ids() == []
    with db.begin() as connection:
        connection.execute(
            update(Document.__table__).values(
                created_at=datetime.now(UTC) - timedelta(minutes=5)
            )
        )
    assert worker.stale_document_ids() == [lost]
