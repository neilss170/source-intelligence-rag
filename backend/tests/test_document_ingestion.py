from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from app.models.document import DocumentStatus
from app.services.document_ingestion import (
    DocumentIngestionError,
    DocumentIngestionService,
)
from app.services.uploaded_evidence_repository import UploadedEvidenceRepository


def test_ingests_text_document_and_detects_duplicate(tmp_path: Path) -> None:
    repository = UploadedEvidenceRepository(tmp_path / "index" / "evidence.json")
    service = DocumentIngestionService(tmp_path / "uploads", repository)
    content = (
        b"Returns Policy\n\nCustomers can return unused products within "
        b"30 days with proof of purchase."
    )

    first = service.ingest("returns policy.txt", content)
    duplicate = service.ingest("returns policy.txt", content)

    assert first.status is DocumentStatus.READY
    assert first.evidence_count == 1
    assert duplicate.status is DocumentStatus.DUPLICATE
    assert repository.load()[0].source.filename == "returns_policy.txt"


def test_rejects_unsupported_document(tmp_path: Path) -> None:
    service = DocumentIngestionService(
        tmp_path / "uploads",
        UploadedEvidenceRepository(tmp_path / "index" / "evidence.json"),
    )

    with pytest.raises(DocumentIngestionError, match="Supported formats"):
        service.ingest("archive.zip", b"data")


def test_rejects_empty_document(tmp_path: Path) -> None:
    service = DocumentIngestionService(
        tmp_path / "uploads",
        UploadedEvidenceRepository(tmp_path / "index" / "evidence.json"),
    )

    with pytest.raises(DocumentIngestionError, match="empty"):
        service.ingest("empty.txt", b"")


def test_deletes_document_evidence_and_stored_file(tmp_path: Path) -> None:
    repository = UploadedEvidenceRepository(tmp_path / "index" / "evidence.json")
    uploads_dir = tmp_path / "uploads"
    service = DocumentIngestionService(uploads_dir, repository)
    document = service.ingest("notes.txt", b"A unique fact for deletion.")

    assert service.delete(document.document_id) is True
    assert repository.load() == []
    assert list(uploads_dir.iterdir()) == []
    assert service.delete(document.document_id) is False


def test_ingests_csv_rows_as_structured_evidence(tmp_path: Path) -> None:
    repository = UploadedEvidenceRepository(tmp_path / "index" / "evidence.json")
    service = DocumentIngestionService(tmp_path / "uploads", repository)

    result = service.ingest(
        "products.csv",
        b"product,price,currency\nKeyboard,4999,INR\nMouse,1299,INR\n",
    )

    evidence = repository.load()
    assert result.evidence_count == 2
    assert "product: Keyboard" in evidence[0].text
    assert "price: 4999" in evidence[0].text


def test_ingests_json_records(tmp_path: Path) -> None:
    repository = UploadedEvidenceRepository(tmp_path / "index" / "evidence.json")
    service = DocumentIngestionService(tmp_path / "uploads", repository)

    result = service.ingest(
        "items.json",
        b'[{"name":"Keyboard","stock":4},{"name":"Mouse","stock":8}]',
    )

    assert result.evidence_count == 2
    assert "Keyboard" in repository.load()[0].text


def test_ingests_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    repository = UploadedEvidenceRepository(tmp_path / "index" / "evidence.json")
    service = DocumentIngestionService(tmp_path / "uploads", repository)
    document = Document()
    document.add_heading("Product guide", level=1)
    document.add_paragraph("The keyboard uses tactile mechanical switches.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Warranty"
    table.rows[0].cells[1].text = "Two years"
    buffer = BytesIO()
    document.save(buffer)

    result = service.ingest("guide.docx", buffer.getvalue())

    assert result.evidence_count == 1
    assert "tactile mechanical switches" in repository.load()[0].text
    assert "Warranty | Two years" in repository.load()[0].text


def test_ingests_xlsx_worksheets(tmp_path: Path) -> None:
    repository = UploadedEvidenceRepository(tmp_path / "index" / "evidence.json")
    service = DocumentIngestionService(tmp_path / "uploads", repository)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["name", "price"])
    sheet.append(["Keyboard", 4999])
    buffer = BytesIO()
    workbook.save(buffer)

    result = service.ingest("products.xlsx", buffer.getvalue())

    assert result.evidence_count == 1
    assert "Worksheet: Products" in repository.load()[0].text
    assert "name: Keyboard" in repository.load()[0].text
