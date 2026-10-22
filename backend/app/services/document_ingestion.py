import csv
import io
import json
import re
from hashlib import sha256
from pathlib import Path

from app.models.document import DocumentResponse, DocumentStatus
from app.models.resume import EvidenceUnit, ResumeSection, SourceReference
from app.services.document_extraction import (
    DocumentExtractionError,
    ExtractedPage,
    extract_pdf,
)
from app.services.redaction import redact_contact_details
from app.services.text_chunking import chunk_pages
from app.services.uploaded_evidence_repository import UploadedEvidenceRepository

SUPPORTED_EXTENSIONS = {
    ".csv",
    ".docx",
    ".json",
    ".md",
    ".markdown",
    ".pdf",
    ".txt",
    ".xlsx",
}
SAFE_FILENAME = re.compile(r"[^a-zA-Z0-9._-]+")


class DocumentIngestionError(ValueError):
    """Raised when an uploaded document cannot be indexed."""


class DocumentIngestionService:
    def __init__(
        self,
        uploads_dir: Path,
        repository: UploadedEvidenceRepository,
        *,
        max_file_size_mb: int = 10,
        redact: bool = True,
        ocr: bool = False,
    ) -> None:
        self._uploads_dir = uploads_dir
        self._repository = repository
        self._max_bytes = max_file_size_mb * 1024 * 1024
        self._redact = redact
        self._ocr = ocr

    def ingest(self, filename: str, content: bytes) -> DocumentResponse:
        safe_name = _safe_filename(filename)
        extension = Path(safe_name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise DocumentIngestionError(
                "Supported formats are PDF, TXT, Markdown, CSV, JSON, DOCX, and XLSX."
            )
        if not content:
            raise DocumentIngestionError("The uploaded document is empty.")
        if len(content) > self._max_bytes:
            raise DocumentIngestionError("The document exceeds the upload size limit.")

        content_hash = sha256(content).hexdigest()
        document_id = f"doc_{content_hash[:12]}"
        existing = self._repository.load()
        existing_for_document = [
            item for item in existing if item.source.document_id == document_id
        ]
        if existing_for_document:
            return DocumentResponse(
                document_id=document_id,
                filename=safe_name,
                status=DocumentStatus.DUPLICATE,
                evidence_count=len(existing_for_document),
            )

        self._uploads_dir.mkdir(parents=True, exist_ok=True)
        stored_path = self._uploads_dir / f"{document_id}_{safe_name}"
        stored_path.write_bytes(content)
        try:
            pages = _extract_pages(
                stored_path, extension, redact=self._redact, ocr=self._ocr
            )
            # A PDF paragraph can run across a page break; a CSV row or JSON
            # record is a page of its own and must stay one.
            chunks = chunk_pages(pages, merge_pages=extension == ".pdf")
            if not chunks:
                raise DocumentIngestionError("The document contains no usable text.")
            additions = [
                EvidenceUnit(
                    evidence_id=f"ev_upload_{content_hash[:12]}_{index:03d}",
                    section=ResumeSection.DOCUMENT,
                    title=_title(Path(safe_name).stem, chunk.heading),
                    text=chunk.text,
                    source=SourceReference(
                        document_id=document_id,
                        filename=safe_name,
                        page=chunk.page,
                        source_text=chunk.text,
                    ),
                )
                for index, chunk in enumerate(chunks, start=1)
            ]
            self._repository.save([*existing, *additions])
        except DocumentIngestionError:
            stored_path.unlink(missing_ok=True)
            raise
        except (
            DocumentExtractionError,
            UnicodeDecodeError,
            json.JSONDecodeError,
            csv.Error,
            ValueError,
            OSError,
        ) as exc:
            stored_path.unlink(missing_ok=True)
            raise DocumentIngestionError(str(exc)) from exc

        return DocumentResponse(
            document_id=document_id,
            filename=safe_name,
            status=DocumentStatus.READY,
            evidence_count=len(additions),
        )

    def delete(self, document_id: str) -> bool:
        if not re.fullmatch(r"doc_[a-f0-9]{12}", document_id):
            return False

        removed_count = self._repository.delete_document(document_id)
        if not removed_count:
            return False

        if self._uploads_dir.exists():
            for stored_path in self._uploads_dir.glob(f"{document_id}_*"):
                if stored_path.is_file():
                    stored_path.unlink(missing_ok=True)
        return True


def _title(stem: str, heading: str) -> str:
    """Passage title: the file, then the section path inside it."""
    return f"{stem} > {heading}" if heading else stem


def _extract_pages(
    path: Path, extension: str, *, redact: bool = True, ocr: bool = False
) -> tuple[ExtractedPage, ...]:
    if extension == ".pdf":
        return extract_pdf(path, redact=redact, ocr=ocr).pages
    if extension in {".txt", ".md", ".markdown"}:
        pages = (_text_page(path.read_text(encoding="utf-8"), 1),)
    elif extension == ".csv":
        pages = _extract_csv(path)
    elif extension == ".json":
        pages = _extract_json(path)
    elif extension == ".docx":
        pages = _extract_docx(path)
    elif extension == ".xlsx":
        pages = _extract_xlsx(path)
    else:
        raise DocumentIngestionError("The document format is not supported.")
    if redact:
        pages = tuple(
            ExtractedPage(page=page.page, text=redact_contact_details(page.text))
            for page in pages
        )
    return pages


def _text_page(text: str, page: int) -> ExtractedPage:
    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    return ExtractedPage(page=page, text=normalized)


def _extract_csv(path: Path) -> tuple[ExtractedPage, ...]:
    text = path.read_text(encoding="utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ()
    headers = [
        value.strip() or f"column_{index + 1}" for index, value in enumerate(rows[0])
    ]
    pages = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = []
        for index, value in enumerate(row):
            if not value.strip():
                continue
            header = headers[index] if index < len(headers) else f"column_{index + 1}"
            values.append(f"{header}: {value.strip()}")
        if values:
            pages.append(
                _text_page(
                    f"CSV row {row_number}\n" + "\n".join(values), row_number - 1
                )
            )
    return tuple(pages)


def _extract_json(path: Path) -> tuple[ExtractedPage, ...]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    records = payload if isinstance(payload, list) else [payload]
    return tuple(
        _text_page(
            f"JSON record {index}\n" + json.dumps(record, ensure_ascii=False, indent=2),
            index,
        )
        for index, record in enumerate(records, start=1)
        if record not in (None, "", [], {})
    )


def _extract_docx(path: Path) -> tuple[ExtractedPage, ...]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    # Walk the body in order so a table stays next to the text that
    # introduces it, and Word headings become section headings.
    blocks: list[str] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _heading_level(paragraph.style.name if paragraph.style else "")
            blocks.append(f"{'#' * level} {text}" if level else text)
        elif tag == "tbl":
            rows = []
            for row in Table(element, document).rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                blocks.append("\n".join(rows))
    return (_text_page("\n\n".join(blocks), 1),) if blocks else ()


def _heading_level(style_name: str) -> int:
    if style_name == "Title":
        return 1
    match = re.fullmatch(r"Heading (\d)", style_name)
    return int(match.group(1)) if match else 0


def _extract_xlsx(path: Path) -> tuple[ExtractedPage, ...]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    pages: list[ExtractedPage] = []
    try:
        for sheet_number, sheet in enumerate(workbook.worksheets, start=1):
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [
                str(value or f"column_{index + 1}")
                for index, value in enumerate(rows[0])
            ]
            lines = [f"## Worksheet: {sheet.title}", ""]
            for row_number, row in enumerate(rows[1:], start=2):
                values = [
                    f"{headers[index]}: {value}"
                    for index, value in enumerate(row)
                    if value not in (None, "")
                ]
                if values:
                    lines.append(f"Row {row_number}: " + " | ".join(values))
            if len(lines) > 2:
                pages.append(_text_page("\n".join(lines), sheet_number))
    finally:
        workbook.close()
    return tuple(pages)


def _safe_filename(filename: str) -> str:
    name = Path(filename or "document").name
    cleaned = SAFE_FILENAME.sub("_", name).strip("._")
    return cleaned or "document"


# Public names for the background ingestion pipeline (app.ingest), which
# reuses the measured extraction and titling rules above unchanged.
extract_pages = _extract_pages
passage_title = _title
safe_filename = _safe_filename
