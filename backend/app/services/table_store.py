"""Uploaded spreadsheets as SQL tables, so totals are computed, not guessed.

Retrieval hands the model a few rows; "what was the total cost?" needs all of
them. So every uploaded CSV and XLSX sheet is also loaded into an in-memory
SQLite database, and the model may ask for one ``SELECT`` whose result it then
cites like any other passage.

The model writes the SQL, so the database is locked down: an authorizer
allows reading and nothing else, ``query_only`` is on, a progress handler
aborts anything that runs too long, and only one statement is accepted.
SQLite ships with Python, so none of this adds a dependency.
"""

import csv
import re
import sqlite3
import threading
import time
from dataclasses import dataclass
from pathlib import Path

#: Rows of a result shown to the model.
MAX_RESULT_ROWS = 50

#: Wall-clock limit for one query.
QUERY_TIMEOUT_S = 2.0

_ALLOWED_ACTIONS = {
    sqlite3.SQLITE_SELECT,
    sqlite3.SQLITE_READ,
    sqlite3.SQLITE_FUNCTION,
    getattr(sqlite3, "SQLITE_RECURSIVE", 33),
}
_IDENTIFIER = re.compile(r"[^a-z0-9_]+")


class TableQueryError(ValueError):
    """Raised when a query is refused or fails; the message goes back to the model."""


@dataclass(frozen=True, slots=True)
class Table:
    name: str
    filename: str
    document_id: str
    columns: tuple[tuple[str, str], ...]
    row_count: int
    sample: tuple[tuple, ...]


@dataclass(frozen=True, slots=True)
class QueryResult:
    sql: str
    columns: tuple[str, ...]
    rows: tuple[tuple, ...]
    truncated: bool
    tables: tuple[Table, ...]

    def render(self) -> str:
        header = " | ".join(self.columns)
        lines = [header] + [
            " | ".join(_cell(value) for value in row) for row in self.rows
        ]
        if self.truncated:
            lines.append(f"(first {MAX_RESULT_ROWS} rows shown)")
        if not self.rows:
            lines.append("(no rows)")
        return f"SQL: {self.sql}\n" + "\n".join(lines)


class TableStore:
    def __init__(self, files: list[tuple[str, str, Path]]) -> None:
        """Load ``(document_id, filename, path)`` spreadsheets into SQLite."""
        self._connection = sqlite3.connect(":memory:", check_same_thread=False)
        self._lock = threading.Lock()
        self.tables: list[Table] = []
        for document_id, filename, path in files:
            for sheet, rows in _read(path):
                if len(rows) < 2:
                    continue
                self._load(document_id, filename, sheet, rows)
        self._connection.execute("PRAGMA query_only = ON")
        self._connection.set_authorizer(_authorize)

    def __bool__(self) -> bool:
        return bool(self.tables)

    def describe(self) -> str:
        """Schemas and a few sample rows, for the model to write SQL against."""
        blocks = []
        for table in self.tables:
            columns = ", ".join(f'"{name}" {kind}' for name, kind in table.columns)
            sample = "\n".join(
                "  " + " | ".join(_cell(value) for value in row) for row in table.sample
            )
            blocks.append(
                f'"{table.name}" from {table.filename}: {table.row_count} rows\n'
                f"  columns: {columns}\n  sample rows:\n{sample}"
            )
        return "\n\n".join(blocks)

    def query(self, sql: str) -> QueryResult:
        sql = sql.strip().rstrip(";").strip()
        if not re.match(r"(?is)^\s*(select|with)\b", sql):
            raise TableQueryError("Only a single SELECT statement is allowed.")
        deadline = time.monotonic() + QUERY_TIMEOUT_S
        with self._lock:
            self._connection.set_progress_handler(
                lambda: int(time.monotonic() > deadline), 10_000
            )
            try:
                cursor = self._connection.execute(sql)
                rows = cursor.fetchmany(MAX_RESULT_ROWS + 1)
                columns = tuple(d[0] for d in cursor.description or ())
            except sqlite3.Error as exc:
                raise TableQueryError(f"SQLite error: {exc}") from exc
            finally:
                self._connection.set_progress_handler(None, 0)
        used = tuple(t for t in self.tables if re.search(rf'\b"?{t.name}"?\b', sql))
        return QueryResult(
            sql=sql,
            columns=columns,
            rows=tuple(rows[:MAX_RESULT_ROWS]),
            truncated=len(rows) > MAX_RESULT_ROWS,
            tables=used or tuple(self.tables[:1]),
        )

    def _load(self, document_id: str, filename: str, sheet: str, rows: list) -> None:
        headers = _unique(
            [_name(str(h or ""), f"column_{i + 1}") for i, h in enumerate(rows[0])]
        )
        body = [list(row) + [None] * (len(headers) - len(row)) for row in rows[1:]]
        body = [
            row[: len(headers)] for row in body if any(v not in (None, "") for v in row)
        ]
        kinds = [_infer([row[i] for row in body]) for i in range(len(headers))]
        typed = [
            tuple(_convert(value, kind) for value, kind in zip(row, kinds, strict=True))
            for row in body
        ]
        stem = _name(Path(filename).stem, "table")
        name = _unique_against(
            f"{stem}_{_name(sheet, 'sheet')}" if sheet else stem,
            {t.name for t in self.tables},
        )
        columns = ", ".join(f'"{h}" {k}' for h, k in zip(headers, kinds, strict=True))
        self._connection.execute(f'CREATE TABLE "{name}" ({columns})')
        placeholders = ", ".join("?" for _ in headers)
        self._connection.executemany(
            f'INSERT INTO "{name}" VALUES ({placeholders})', typed
        )
        self.tables.append(
            Table(
                name=name,
                filename=filename,
                document_id=document_id,
                columns=tuple(zip(headers, kinds, strict=True)),
                row_count=len(typed),
                sample=tuple(typed[:3]),
            )
        )


def _authorize(action, *_args) -> int:
    return sqlite3.SQLITE_OK if action in _ALLOWED_ACTIONS else sqlite3.SQLITE_DENY


def _read(path: Path) -> list[tuple[str, list]]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return [("", list(csv.reader(handle)))]
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()
    return []


def _infer(values: list) -> str:
    present = [v for v in values if v not in (None, "")]
    if not present:
        return "TEXT"
    if all(_number(v) is not None and float(_number(v)).is_integer() for v in present):
        return "INTEGER"
    if all(_number(v) is not None for v in present):
        return "REAL"
    return "TEXT"


def _number(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _convert(value, kind: str):
    if value in (None, ""):
        return None
    if kind == "INTEGER":
        return int(_number(value))
    if kind == "REAL":
        return _number(value)
    return str(value)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    return str(value)


def _name(text: str, fallback: str) -> str:
    cleaned = _IDENTIFIER.sub("_", text.lower()).strip("_")
    if not cleaned:
        return fallback
    return f"t_{cleaned}" if cleaned[0].isdigit() else cleaned


def _unique(names: list[str]) -> list[str]:
    seen: set[str] = set()
    return [_unique_against(name, seen, add=True) for name in names]


def _unique_against(name: str, taken: set[str], *, add: bool = False) -> str:
    candidate, n = name, 2
    while candidate in taken:
        candidate, n = f"{name}_{n}", n + 1
    if add:
        taken.add(candidate)
    return candidate
